"""Publish this week's uploads while preserving all prior content."""
import copy
import hashlib
import html
import importlib.util
import json
from pathlib import Path
import re

from docx import Document
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph

SPEC = importlib.util.spec_from_file_location("weekly", Path(__file__).with_name("publish-oct12-18-2026.py"))
weekly = importlib.util.module_from_spec(SPEC)
SPEC.loader.exec_module(weekly)
weekly.DATE = "2026-10-30"
SITE, UPLOADS = weekly.SITE, weekly.UPLOADS
WEEK = "19–25 October 2026"
BLOG = "/blogs/healthcare-market-entry-india-hospital-adoption/"
CASE = "/case-studies/medical-device-india-hospital-adoption/"
NEWS = [
    ("healthcare", "healthcare-brief", "GreyRadius_Healthcare", "Mercury-device phase-out, rural healthcare funding and FDA decisions",
     "India 2030 target · US nursing workforce · Rural access and drug approvals"),
    ("charged", "ev-battery-intelligence", "The_Charged", "India battery storage, regional EV demand and solid-state manufacturing",
     "Failte–Sigenergy · ACME and SECI projects · Europe BEV registrations"),
    ("the-signal", "enterprise-ai-intelligence", "The_Signal", "Enterprise AI adoption, governance and execution priorities",
     "Enterprise platforms · Agent deployment · Infrastructure and operating controls"),
    ("the-stack", "infrastructure-brief", "The_Stack", "GPU deployments, sovereign cloud and AI power infrastructure",
     "Noida and Chennai compute · Saudi and UAE capacity · Power flexibility"),
    ("the-wafer", "semiconductor-intelligence", "The_Wafer", "AI packaging capacity, equipment earnings and semiconductor materials",
     "NVIDIA supply chain · Lam Research · Texas Instruments · CMP technology"),
]


def inline(paragraph):
    """Keep run emphasis and DOCX hyperlink relationships, never Word XML."""
    def run(element):
        text = "".join(html.escape(t.text or "") for t in element.iter(qn("w:t")))
        props = element.find(qn("w:rPr"))
        if props is not None:
            if props.find(qn("w:b")) is not None:
                text = "<strong>" + text + "</strong>"
            if props.find(qn("w:i")) is not None:
                text = "<em>" + text + "</em>"
        return text
    parts = []
    for element in paragraph._p:
        if element.tag == qn("w:r"):
            parts.append(run(element))
        elif element.tag == qn("w:hyperlink"):
            label = "".join(run(r) for r in element.findall(qn("w:r")))
            rid = element.get(qn("r:id"))
            url = paragraph.part.rels[rid].target_ref if rid else "#" + element.get(qn("w:anchor"), "")
            parts.append(f'<a href="{html.escape(url, quote=True)}">{label}</a>')
    return "".join(parts)


def convert_docx(source, skip):
    doc = Document(source)
    chunks, position, listing = [], 0, False
    for element in doc.element.body:
        if element.tag == qn("w:p"):
            paragraph = Paragraph(element, doc)
            position += 1
            if position <= skip or not paragraph.text.strip():
                continue
            bullet = paragraph.style.name.startswith("List")
            if bullet and not listing:
                chunks.append("<ul>")
                listing = True
            if listing and not bullet:
                chunks.append("</ul>")
                listing = False
            style = paragraph.style.name
            tag = "li" if bullet else "h2" if style == "Heading 1" else "h3" if style.startswith("Heading") else "p"
            chunks.append(f"<{tag}>{inline(paragraph)}</{tag}>")
        elif element.tag == qn("w:tbl"):
            if listing:
                chunks.append("</ul>")
                listing = False
            table = Table(element, doc)
            chunks.append("<table><thead>")
            for i, row in enumerate(table.rows):
                if i == 1:
                    chunks.append("</thead><tbody>")
                tag = "th" if i == 0 else "td"
                chunks.append("<tr>" + "".join(
                    f"<{tag}>" + "<br/>".join(inline(p) for p in cell.paragraphs) + f"</{tag}>"
                    for cell in row.cells
                ) + "</tr>")
            chunks.append("</tbody></table>")
    if listing:
        chunks.append("</ul>")
    return doc, "\n".join(chunks)


def main():
    old = {p: hashlib.sha256(p.read_bytes()).hexdigest()
           for section in ["insights/newsletters", "blogs", "case-studies"]
           for p in (SITE / section).glob("**/index.html")
           if len(p.relative_to(SITE / section).parts) > (2 if section == "insights/newsletters" else 1)
           and "oct19-25-2026" not in str(p) and BLOG.strip("/") not in str(p) and CASE.strip("/") not in str(p)}
    latest = (SITE / "assets/js/newsletter-data.js").read_text()
    hub = weekly.parse((SITE / "insights/newsletters/index.html").read_text())
    urls = []
    for series, prefix, source_prefix, heading, summary in NEWS:
        source = next(UPLOADS.glob(source_prefix + "*179335311*.html"))
        raw = source.read_text()
        original = weekly.parse(raw)
        url = f"/insights/newsletters/{series}/{prefix}-oct19-25-2026/"
        navigation = (
            f'<nav aria-label="Newsletter navigation" style="padding:.85rem 1rem;background:#132846;color:white;font:14px sans-serif;">'
            f'<a href="/insights/newsletters/{series}/" style="color:white;">All editions</a> · '
            '<a href="/insights/newsletters/" style="color:white;">Newsletter hub</a></nav>'
        )
        if series == "the-stack":
            navigation += (
                '<aside data-editorial-note="continuing-coverage" style="max-width:920px;margin:1rem auto;padding:1rem;background:#fff6e8;color:#132846;font:14px/1.6 sans-serif;">'
                '<strong>Coverage note:</strong> The Microsoft/Nscale item revisits a development covered in the '
                '<a href="/insights/newsletters/the-stack/infrastructure-brief-oct12-18-2026/">12–18 October edition</a>. '
                'It is retained as continuing coverage, not a separate new project.</aside>'
            )
        result = raw.replace("</head>", weekly.metadata(original.title.get_text(), heading + ". " + summary + ".", url) + "</head>", 1)
        result = re.sub(r"(<body\b[^>]*>)", lambda m: m[0] + navigation, result, count=1)
        check = weekly.parse(result)
        for extra in check.select('nav[aria-label="Newsletter navigation"],aside[data-editorial-note]'):
            extra.decompose()
        assert str(check.body) == str(original.body), source.name
        weekly.save(url.lstrip("/") + "index.html", result)
        archive = weekly.parse((SITE / f"insights/newsletters/{series}/index.html").read_text())
        if not archive.find("a", href=url):
            current = archive.select_one("a[data-newsletter-current]")
            card = copy.deepcopy(current)
            for tag in [current, *current.select("[data-newsletter-issue],[data-newsletter-headline]")]:
                for attr in ["data-newsletter-current", "data-newsletter-issue", "data-newsletter-headline"]:
                    tag.attrs.pop(attr, None)
            card["href"] = url
            card.select_one("[data-newsletter-issue]").string = "Weekly edition · " + WEEK
            card.select_one("[data-newsletter-headline]").string = heading
            card.find_all(["p", "h2"])[-1].string = summary
            current.insert_before(card)
        weekly.save(f"insights/newsletters/{series}/index.html", str(archive))
        link = hub.select_one(f'[data-newsletter-latest="{series}"]')
        link["href"] = url
        link.string = re.sub(r" for .*$", " for " + WEEK, link.get_text())
        def update(match):
            record = re.sub(r'url: "[^"]+"', 'url: "' + url + '"', match[2])
            record = re.sub(r'date: "[^"]+"', 'date: "' + WEEK + '"', record)
            record = re.sub(r'headline:\s*"[^"]+"', "headline: " + json.dumps(heading), record)
            # This upload identifies editions by dates; do not invent issue numbers.
            record = re.sub(r'\n\s*issue: "\d+",', "", record)
            return match[1] + record + match[3]
        latest, count = re.subn(rf'("?{re.escape(series)}"?: Object\.freeze\(\{{)([\s\S]*?)(\n    \}}\))', update, latest, count=1)
        assert count == 1
        urls.append(url)
    weekly.save("assets/js/newsletter-data.js", latest)
    weekly.save("insights/newsletters/index.html", str(hub))

    blog_doc, blog_body = convert_docx(next(UPLOADS.glob("Healthcare_Market_Entry_India*179335311*.docx")), 1)
    case_doc, case_body = convert_docx(next(UPLOADS.glob("GreyRadius_Medical_Technology_Case_Study*179335311*.docx")), 2)
    title, case_title = blog_doc.paragraphs[0].text, case_doc.paragraphs[0].text
    weekly.make_article(
        "blogs/gtm-execution-as-a-service-embedded-team/index.html", BLOG, title,
        "Connect medical-device approval with hospital adoption in India: clinical evidence, procurement, budgets, installation, training and service readiness.",
        "Healthcare · India Market Entry", blog_body,
        related=(CASE, "Read the related Indian hospital adoption case study"),
    )
    weekly.make_article(
        "case-studies/market-entry-qualified-pipeline-gtm-execution/index.html", CASE, case_title,
        case_doc.paragraphs[1].text, "MedTech · India Market Entry", case_body, case_doc.paragraphs[1].text,
        related=(BLOG, "Explore the healthcare market-entry guide"),
    )
    for index, previous, url, title in [
        ("insights", "/blogs/gtm-execution-as-a-service-embedded-team/", BLOG, title),
        ("case-studies", "/case-studies/market-entry-qualified-pipeline-gtm-execution/", CASE, case_title),
    ]:
        page = weekly.parse((SITE / index / "index.html").read_text())
        if not page.find("a", href=url):
            previous_card = page.find("a", href=previous)
            card = copy.deepcopy(previous_card)
            card["href"] = url
            card.find(["h2", "h3"]).string = title
            if index == "insights":
                card["data-industry"] = "healthcare"
                card.select_one(".cs-photo-tag").string = "Healthcare · India"
                card.select_one(".case-study-card__tag").string = "Healthcare / Market Entry"
                card.select_one(".case-study-card__image div span").string = "MedTech"
            else:
                card["data-industry"], card["data-geo"] = "healthcare", "india"
                card.select_one(".cs-card-image span").string = "Indian Hospital Adoption"
                card.select_one(".tag-geo").string = "India"
                card.select_one(".tag-industry").string = "Healthcare"
                card.select_one(".cs-card-outcome").string = "Aligned clinical evaluation, purchasing authority, hospital economics and service readiness. Conversion and revenue outcomes remain unverified."
                card.select_one(".cs-card-kpi").string = "Adoption gates · hospital economics · local support"
            previous_card.insert_before(card)
        weekly.save(index + "/index.html", str(page))
    sitemap = (SITE / "sitemap.xml").read_text()
    entries = "".join(
        f"<url><loc>https://greyradius.com{url}</loc><lastmod>{weekly.DATE}</lastmod><changefreq>monthly</changefreq><priority>0.7</priority></url>\n"
        for url in urls + [BLOG, CASE] if "https://greyradius.com" + url + "</loc>" not in sitemap
    )
    weekly.save("sitemap.xml", sitemap.replace("</urlset>", entries + "</urlset>"))
    for p, digest in old.items():
        assert hashlib.sha256(p.read_bytes()).hexdigest() == digest, f"Changed older content: {p}"
    print(f"Added seven pages. Preserved {len(old)} prior content pages byte-for-byte.")


if __name__ == "__main__":
    main()
